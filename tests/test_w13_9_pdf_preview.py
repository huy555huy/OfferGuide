"""W13.9 — docx HTML preview + PDF endpoint tests."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import offerguide
from offerguide.config import Settings
from offerguide.profile import UserProfile
from offerguide.skills import SkillRuntime, discover_skills
from offerguide.ui.notify import ConsoleNotifier
from offerguide.ui.web import create_app

SKILLS_ROOT = Path(__file__).parent.parent / "src/offerguide/skills"


# ═══════════════════════════════════════════════════════════════════
# preview module
# ═══════════════════════════════════════════════════════════════════


class TestDocxPreviewModule:
    def test_missing_file_raises(self):
        from offerguide.skills.tailor_resume.preview import docx_to_html
        with pytest.raises(FileNotFoundError):
            docx_to_html("/nonexistent/path.docx")

    def test_real_docx_renders_to_html(self, tmp_path):
        """Build a tiny docx in-memory + verify mammoth produces HTML."""
        from offerguide.skills.tailor_resume.preview import docx_to_html
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        # Create minimal docx
        doc = Document()
        doc.add_heading("简历预览测试", 0)
        doc.add_paragraph("这是一段测试段落文本。")
        doc.add_heading("教育背景", 1)
        doc.add_paragraph("某高校 · 应用统计专硕 · 2025-2027")
        docx_path = tmp_path / "test.docx"
        doc.save(str(docx_path))

        html = docx_to_html(docx_path)
        # Banner + at least one of our content strings
        assert "简历预览测试" in html
        assert "应用统计" in html
        # Print-optimized CSS present
        assert "@page" in html
        assert "<style>" in html.lower()

    def test_soffice_returns_none_when_not_installed(self, tmp_path, monkeypatch):
        """If neither soffice nor libreoffice is on PATH and Mac LibreOffice.app
        isn't present, the function should return None (not raise)."""
        from offerguide.skills.tailor_resume import preview

        monkeypatch.setattr(shutil, "which", lambda _: None)
        monkeypatch.setattr(preview, "_try_mac_libreoffice", lambda: None)

        # Even without docx file, returns None first (since soffice missing).
        # Make sure the module returns None without raising.
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")
        doc = Document()
        doc.add_paragraph("x")
        docx_path = tmp_path / "x.docx"
        doc.save(str(docx_path))

        result = preview.soffice_to_pdf(docx_path, tmp_path / "out")
        assert result is None


# ═══════════════════════════════════════════════════════════════════
# Web routes
# ═══════════════════════════════════════════════════════════════════


@pytest.fixture
def app_client(tmp_path, monkeypatch):
    # CD into tmp_path so the route's "data/tailored/<filename>" lookups land here
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data" / "tailored").mkdir(parents=True)

    store = offerguide.Store(tmp_path / "preview.db")
    store.init_schema()
    skills = discover_skills(SKILLS_ROOT)
    s = Settings(deepseek_api_key="x", default_model="stub")

    from offerguide.llm import LLMResponse
    class _StubLLM:
        def chat(self, messages, **kw):
            return LLMResponse(content="{}", model="stub")
    runtime = SkillRuntime(llm=_StubLLM(), store=store)
    profile = UserProfile(raw_resume_text="x", source_pdf="/tmp/x.pdf")
    app = create_app(
        settings=s, store=store, profile=profile,
        skills=skills, runtime=runtime, notifier=ConsoleNotifier(),
    )
    return TestClient(app), tmp_path


class TestPreviewRoute:
    def _make_real_docx(self, tmp_path: Path, filename: str):
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")
        doc = Document()
        doc.add_heading("Real Tailored Resume", 0)
        doc.add_paragraph("Some content here for preview.")
        path = tmp_path / "data" / "tailored" / filename
        doc.save(str(path))
        return path

    def test_preview_404_when_file_missing(self, app_client):
        client, _ = app_client
        resp = client.get("/api/tailor/preview/missing.docx")
        assert resp.status_code == 404

    def test_preview_rejects_path_traversal(self, app_client):
        client, _ = app_client
        for evil in ("../etc/passwd", "..\\\\windows", "subdir/foo.docx"):
            resp = client.get(f"/api/tailor/preview/{evil}")
            assert resp.status_code in (400, 404)

    def test_preview_rejects_non_docx(self, app_client):
        client, _ = app_client
        resp = client.get("/api/tailor/preview/file.txt")
        assert resp.status_code == 400

    def test_preview_returns_html(self, app_client):
        client, tmp_path = app_client
        self._make_real_docx(tmp_path, "tailored_test.docx")
        resp = client.get("/api/tailor/preview/tailored_test.docx")
        assert resp.status_code == 200
        # Content-Type is HTML
        assert "text/html" in resp.headers["content-type"]
        # Renders our content
        assert "Real Tailored Resume" in resp.text
        assert "Some content here" in resp.text
        # Print CSS present
        assert "@page" in resp.text


class TestPDFRoute:
    def test_pdf_404_when_file_missing(self, app_client):
        client, _ = app_client
        resp = client.get("/api/tailor/pdf/missing.docx")
        assert resp.status_code == 404

    def test_pdf_rejects_traversal(self, app_client):
        client, _ = app_client
        resp = client.get("/api/tailor/pdf/../etc/passwd")
        # FastAPI normalizes path so this resolves to /api/tailor/etc/passwd
        # which doesn't match any route
        assert resp.status_code in (400, 404)

    def test_pdf_503_when_libreoffice_missing(self, app_client, monkeypatch):
        client, tmp_path = app_client
        # Ensure libreoffice not found
        monkeypatch.setattr(shutil, "which", lambda _: None)
        from offerguide.skills.tailor_resume import preview
        monkeypatch.setattr(preview, "_try_mac_libreoffice", lambda: None)

        # Need a real docx file present
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")
        doc = Document()
        doc.add_paragraph("x")
        path = tmp_path / "data" / "tailored" / "test.docx"
        doc.save(str(path))

        resp = client.get("/api/tailor/pdf/test.docx")
        assert resp.status_code == 503
        assert "libreoffice" in resp.text or "LibreOffice" in resp.text or "Save as PDF" in resp.text
